"""Regression coverage for version matching and review conclusions."""
from copy import copy
from io import BytesIO
import unittest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from inspector import inspect
from comparison import compare_reports


def report(workbook, name='test.xlsx'):
    stream=BytesIO();workbook.save(stream)
    return inspect(stream.getvalue(),name)


class ComparisonTest(unittest.TestCase):
    def sample(self):
        book=Workbook();sheet=book.active;sheet.title='明细'
        for row in [('项目','金额'),('设备',100),('服务',200),('合计','=SUM(B2:B3)')]:sheet.append(row)
        return book

    def test_identical_has_no_changes(self):
        data=report(self.sample())
        self.assertEqual(compare_reports(data,data)['summary']['total'],0)

    def test_content_and_format_change(self):
        book=self.sample();old=report(book)
        book.active['B2']=120
        book.active['B2'].font=Font(name='Changed',size=14,bold=True)
        book.active['B2'].alignment=Alignment(horizontal='center',wrap_text=True)
        result=compare_reports(old,report(book))
        change=next(c for c in result['sheets'][0]['changes'] if c['new'] and c['new']['address']=='B2')
        self.assertEqual(change['old']['value'],'100')
        self.assertTrue(change['content_changed']);self.assertTrue(change['review'])
        self.assertTrue({'font','size','horizontal','wrap'} <= {d['field'] for d in change['format_changes']})

    def test_insert_row_tracks_existing_content(self):
        book=self.sample();old=report(book)
        book.active.insert_rows(2);book.active['A2']='新增';book.active['B2']=50
        result=compare_reports(old,report(book))
        changes=result['sheets'][0]['changes']
        moved=next(c for c in changes if c['new'] and c['new']['address']=='A3')
        self.assertEqual(moved['old']['address'],'A2')
        self.assertFalse(moved['content_changed']);self.assertFalse(moved['format_changes'])
        added=next(c for c in changes if c['new'] and c['new']['address']=='B2')
        self.assertEqual(added['kind'],'新增或未匹配');self.assertTrue(added['review'])
        self.assertIsNotNone(added['reference'])

    def test_insert_column_tracks_content(self):
        book=self.sample();old=report(book);book.active.insert_cols(1)
        result=compare_reports(old,report(book))
        change=next(c for c in result['sheets'][0]['changes'] if c['new']['address']=='C2')
        self.assertEqual(change['old']['address'],'B2');self.assertTrue(change['moved'])
        self.assertFalse(change['format_changes'])

    def test_delete_and_move(self):
        book=self.sample();old=report(book);book.active.delete_rows(2)
        result=compare_reports(old,report(book))
        self.assertTrue(any(c['kind']=='删除或未匹配' and c['old']['value']=='设备' for c in result['sheets'][0]['changes']))
        self.assertTrue(any(c['moved'] and c['new'] and c['new']['value']=='服务' for c in result['sheets'][0]['changes']))

    def test_sheet_rename_and_manual_pair(self):
        book=self.sample();old=report(book);book.active.title='新版明细'
        result=compare_reports(old,report(book))
        self.assertEqual(result['summary']['unmatched_sheets'],0)
        self.assertTrue(result['sheets'][0]['review_pairing'])
        book.active['A1']='完全不同的标题'
        result=compare_reports(old,report(book))
        self.assertEqual(result['summary']['unmatched_sheets'],2)
        result=compare_reports(old,report(book),{'新版明细':'明细'})
        self.assertEqual(result['summary']['unmatched_sheets'],0)
        self.assertEqual(result['sheets'][0]['pairing'],'用户指定工作表')

    def test_ambiguous_renames_stay_unmatched(self):
        book=self.sample();book.copy_worksheet(book.active);old=report(book)
        book.worksheets[0].title='新一';book.worksheets[1].title='新二'
        self.assertEqual(compare_reports(old,report(book))['summary']['unmatched_sheets'],4)

    def test_added_sheet_and_metadata(self):
        book=self.sample();old=report(book)
        book.create_sheet('新工作表');book.active.print_area='A1:B4'
        result=compare_reports(old,report(book))
        self.assertEqual(result['summary']['unmatched_sheets'],1)
        self.assertEqual(result['sheets'][0]['metadata'][0]['label'],'打印范围')

    def test_merge_movement_uses_size_not_address(self):
        book=Workbook();sheet=book.active;sheet['A1']='标题';sheet.merge_cells('A1:B1');old=report(book)
        sheet.unmerge_cells('A1:B1');sheet['A1']=None;sheet['A2']='标题';sheet.merge_cells('A2:B2')
        change=next(c for c in compare_reports(old,report(book))['sheets'][0]['changes'] if c['new'] and c['new']['value']=='标题')
        self.assertTrue(change['moved']);self.assertFalse(change['format_changes'])

    def test_same_text_different_type_is_changed(self):
        book=Workbook();book.active['A1']=1;old=report(book);book.active['A1']='1'
        self.assertTrue(compare_reports(old,report(book))['sheets'][0]['changes'][0]['content_changed'])

    def test_invalid_mapping(self):
        data=report(self.sample())
        with self.assertRaises(ValueError):compare_reports(data,data,{'missing':'明细'})

if __name__=='__main__':unittest.main()
