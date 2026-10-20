"""Read-only OOXML inspection, ported from reference/original.vba."""
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import re
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

RULES = ['显示格式','数据有无','水平对齐','垂直对齐','换行与缩小','合并单元格','字体']
COLORS = {'blue':'#b9dfff','green':'#bce8c9','red':'#ffc3c3','yellow':'#ffed9c','none':'#ffffff'}

def color_info(c):
    if c is None: return '未指定'
    return f'{c.type}:{getattr(c,c.type)}' + (f' tint:{c.tint}' if c.tint else '')

def marks(c, font, size, merged=False):
    a=c.alignment; f=c.font; fmt=c.number_format.strip()
    colors=['none']*7
    colors[0]='none' if fmt in ('General','G/標準') else 'blue' if fmt=='@' else 'green' if fmt in ('0_','#,##0_','0.0_') else 'yellow'
    colors[1]='red' if c.data_type=='f' or c.value not in (None,'') else 'none'
    colors[2]={'left':'blue','center':'yellow','right':'green','justify':'red','distributed':'red','fill':'red','centerContinuous':'red'}.get(a.horizontal,'none')
    colors[3]={'top':'blue','center':'yellow','bottom':'green','justify':'red','distributed':'blue'}.get(a.vertical or 'bottom','none')
    colors[4]='blue' if a.wrap_text else 'yellow' if a.shrink_to_fit else 'none'
    colors[5]='blue' if merged else 'none'
    colors[6]='red' if f.name!=font else 'blue' if f.sz==size else 'none'
    # Default Excel theme dark1 is normally black; custom theme mapping remains uncertain.
    black=f.color is None or (f.color.type=='rgb' and f.color.rgb[-6:]=='000000') or (f.color.type=='theme' and f.color.theme==1 and not f.color.tint) or (f.color.type=='indexed' and f.color.indexed in (8,64))
    return colors, bool(f.b or not black)

def inspect(data, name, font='ＭＳ Ｐゴシック', size=10):
    with ZipFile(BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist())>150*1024*1024: raise ValueError('解压后超过 150 MB，请拆分文件')
        shapes=[]
        for p in z.namelist():
            if re.match(r'xl/drawings/drawing\d+\.xml$',p):
                root=ET.fromstring(z.read(p))
                for el in root.iter():
                    tag=el.tag.rsplit('}',1)[-1]
                    if tag in ('pic','sp','graphicFrame','cxnSp','grpSp'):
                        props=next((n for n in el.iter() if n.tag.endswith('}cNvPr')),None)
                        shapes.append({'part':p,'kind':'图片' if tag=='pic' else tag,'name':props.get('name','') if props is not None else ''})
    wb=openpyxl.load_workbook(BytesIO(data),data_only=False,keep_links=False)
    result={'name':name,'title':wb.properties.title or '', 'sheets':[], 'names':[], 'shapes':shapes,'new_type':False,
      'notes':['颜色表示原 VBA 分类，不等同于错误。','公式仅显示表达式，不执行计算。','显示格式使用 OOXML 格式码，NumberFormatLocal 与 Excel 本地化结果可能不同。','默认主题黑色按 dark1 处理；自定义主题、条件格式、行列默认样式及特殊对象需 Excel 对照验证。','打印范围外沿用 VBA：只比较 UsedRange 右下边界，多区域返回不明；不判断左侧或上侧内容。','预览为单元格属性网格，不模拟 Excel 排版；未实现图形视觉预览。']}
    for n in wb.defined_names.values(): result['names'].append({'name':n.name,'ref':n.attr_text,'scope':'工作簿'})
    for ws in wb.worksheets:
        for n in ws.defined_names.values(): result['names'].append({'name':n.name,'ref':n.attr_text,'scope':ws.title})
        # Sparse iteration avoids a huge rectangle from a far-away styled cell.
        if len(ws._cells)>250000: raise ValueError(f'{ws.title} 超过 250000 个已存储单元格，请拆分')
        ranges=list(ws.merged_cells.ranges)
        if len(ranges)>10000: raise ValueError('合并区域过多')
        area=str(ws.print_area) if ws.print_area else ''
        outside='不明'
        if ws.print_area and len(ws._print_area.ranges)==1:
            r=next(iter(ws._print_area.ranges))
            outside='有り' if ws.max_row>r.max_row or ws.max_column>r.max_col else '無し'
        sheet={'name':ws.title,'rows':ws.max_row,'cols':ws.max_column,'print_area':area,'outside':outside,'state':ws.sheet_state,'merges':[str(r) for r in ranges],'cells':[]}
        for c in ws._cells.values():
            if isinstance(c,MergedCell): continue
            if c.value is None and not c.has_style: continue
            merged=next((str(r) for r in ranges if c.coordinate in r),None)
            colors,pattern=marks(c,font,size,bool(merged))
            val='' if c.value is None else str(c.value)
            if re.search(r'@.*@',val): result['new_type']=True
            sheet['cells'].append({'address':c.coordinate,'row':c.row,'col':c.column,'value':val,'colors':colors,'pattern':pattern,'merge':merged,'format':c.number_format,'font':c.font.name,'size':c.font.sz,'bold':bool(c.font.b),'font_color':color_info(c.font.color),'horizontal':c.alignment.horizontal or 'general','vertical':c.alignment.vertical or 'bottom','wrap':bool(c.alignment.wrap_text),'shrink':bool(c.alignment.shrink_to_fit),'formula':c.data_type=='f'})
        result['sheets'].append(sheet)
    wb.close()
    return result
