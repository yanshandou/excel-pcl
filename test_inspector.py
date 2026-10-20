import unittest
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from inspector import inspect
class InspectionTest(unittest.TestCase):
    def test_rules_and_metadata(self):
        w=Workbook();s=w.active;s.title='检查'
        s['A1']='@model@';s['A1'].number_format='@';s['A1'].font=Font(name='Target',size=10,bold=True,color='FF000000')
        s['A1'].alignment=Alignment(horizontal='left',vertical='distributed',wrap_text=True,shrink_to_fit=True)
        s['B2']='=1+1';s['B2'].font=Font(name='Other',size=10);s.merge_cells('C3:D3');s['C3']='合并';s.print_area='A1:C3';s['E5']='越界'
        s['A2']=1;s['A2'].number_format='0_'
        b=BytesIO();w.save(b);raw=b.getvalue();r=inspect(raw,'test.xlsx','Target',10)
        cells={c['address']:c for c in r['sheets'][0]['cells']}
        self.assertEqual(cells['A1']['colors'],['blue','red','blue','blue','blue','none','blue'])
        self.assertTrue(cells['A1']['pattern']);self.assertEqual(cells['B2']['colors'][6],'red');self.assertEqual(cells['C3']['colors'][5],'blue')
        self.assertEqual(cells['A2']['colors'][0],'green');self.assertEqual(r['sheets'][0]['outside'],'有り');self.assertTrue(r['new_type']);self.assertEqual(raw,b.getvalue())
    def test_missing_and_multi_print_area(self):
        w=Workbook();s=w.active;b=BytesIO();w.save(b)
        self.assertEqual(inspect(b.getvalue(),'blank.xlsx')['sheets'][0]['outside'],'不明')
        s.print_area=['A1:B2','D4:E5'];b=BytesIO();w.save(b)
        self.assertEqual(inspect(b.getvalue(),'multi.xlsx')['sheets'][0]['outside'],'不明')
    def test_input_limits(self):
        for size in (0,410,float('nan'),float('inf')):
            with self.subTest(size=size),self.assertRaisesRegex(ValueError,'有效字体和字号'):
                inspect(b'', 'test.xlsx', 'Target', size)
        with self.assertRaisesRegex(ValueError,'30 MB'):
            inspect(bytes(30*1024*1024+1),'large.xlsx')

    def test_preflight_rejects_large_sheet_and_merges(self):
        from zipfile import ZipFile
        for xml,message in [
            ('<worksheet><sheetData>'+('<c/>'*250001)+'</sheetData></worksheet>','已存储单元格'),
            ('<worksheet><mergeCells><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>','合并区域合计'),
            ('<worksheet><mergeCells>'+('<mergeCell ref="A1:B1"/>'*10001)+'</mergeCells></worksheet>','合并区域过多')]:
            b=BytesIO()
            with ZipFile(b,'w') as z:z.writestr('xl/worksheets/sheet1.xml',xml)
            with self.subTest(message=message),self.assertRaisesRegex(ValueError,message):
                inspect(b.getvalue(),'test.xlsx')

    def test_merge_lookup_across_rows(self):
        w=Workbook();s=w.active
        s.merge_cells('B2:C4');s.merge_cells('D1:E2');s.merge_cells('A6:B7')
        for address in ('B2','D1','A6','A3','F8'):s[address]=address
        b=BytesIO();w.save(b)
        cells={c['address']:c for c in inspect(b.getvalue(),'test.xlsx')['sheets'][0]['cells']}
        for address,merged in [('B2','B2:C4'),('D1','D1:E2'),('A6','A6:B7'),('A3',None),('F8',None)]:
            self.assertEqual(cells[address]['merge'],merged)

if __name__=='__main__':unittest.main()
